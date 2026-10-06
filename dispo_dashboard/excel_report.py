"""
Erzeugt den Excel-Report mit echten Formeln (Stichtag / Horizonte auf "Übersicht" änderbar).
"""
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config

ARIAL = "Arial"
F_NORM = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=ARIAL, size=14, bold=True)
F_INPUT = Font(name=ARIAL, size=10, color="0000FF", bold=True)
F_GREY = Font(name=ARIAL, size=9, italic=True, color="666666")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_RED = PatternFill("solid", fgColor="F8CBAD")
FILL_ORA = PatternFill("solid", fgColor="FCD5B4")
FILL_YEL = PatternFill("solid", fgColor="FFE699")
FILL_GRN = PatternFill("solid", fgColor="C6EFCE")
FILL_FORMULA = PatternFill("solid", fgColor="EEF3FA")
FILL_INFO = PatternFill("solid", fgColor="EDEDED")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
DATUM = "DD.MM.YYYY"
MENGE = "#,##0;-#,##0;-"

ST = "Übersicht!$C$4"
ROT = "Übersicht!$C$5"
GELB = "Übersicht!$C$6"

STATUS_FARBEN = [
    ("Rückstand ungedeckt", FILL_RED), ("Engpass", FILL_RED),
    ("Rückstand, Zugang kommt", FILL_ORA), ("Zugang zu spät", FILL_ORA),
    ("Rückstand gedeckt", FILL_YEL), ("Fehlmenge später", FILL_YEL),
    ("Gedeckt", FILL_GRN),
]


def _kopf(ws, row, namen, breiten=None):
    for i, n in enumerate(namen, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if breiten:
            ws.column_dimensions[get_column_letter(i)].width = breiten[i - 1]
    ws.row_dimensions[row].height = 30


def _wert(v):
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.to_pydatetime()
    return v


def _statusfarben(ws, rng):
    for wert, fill in STATUS_FARBEN:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{wert}"'], fill=fill))


def erstelle(bedarf, bestand_artikel, zugang, artikel, stichtag, ziel, datenstand=None, stockreport=None):
    """
    bedarf: loaders.lade_bedarf
    bestand_artikel: loaders.lade_bestand()[1]
    zugang: loaders.lade_zugang
    artikel: logic.berechne()[1] – nur für Reihenfolge/Kundenliste genutzt
    datenstand: Liste von (Quelle, Datum, Dateiname) – wird auf der Übersicht angezeigt
    stockreport: loaders.lade_stockreport (optional, nur Nebeninfo)
    """
    datenstand = datenstand or []
    wb = Workbook()
    ue = wb.active
    ue.title = "Übersicht"
    pos = wb.create_sheet("Positionen")
    dek = wb.create_sheet("Deckung je Artikel")
    kw = wb.create_sheet("Bedarf je Woche")
    rk = wb.create_sheet("Rückstand")
    bs = wb.create_sheet("Bestand")
    zg = wb.create_sheet("Zugänge")
    sr = wb.create_sheet("Stock Report (Info)") if stockreport is not None else None
    ko = wb.create_sheet("Werke & Kontakte")

    # ---------------- Stock Report (nur Info) ----------------
    SR = None
    if sr is not None:
        sr["A1"] = ("Stock Report des Versenders – NUR NEBENINFO, geht nicht in die Deckungsrechnung ein. "
                    "„Letzte Notiz“ = jeweils neueste nicht leere Notizspalte.")
        sr["A1"].font = F_GREY
        snamen = list(config.STOCKREPORT_SPALTEN.values()) + ["Letzte Notiz", "Notiz vom"]
        _kopf(sr, 2, snamen, [12, 30, 18, 11, 12, 13, 12, 11, 11, 11, 11, 11, 11, 14, 18, 70, 10])
        srlast = max(len(stockreport) + 2, 3)
        for i, rec in enumerate(stockreport[snamen].itertuples(index=False), start=3):
            for j, v in enumerate(rec, 1):
                c = sr.cell(row=i, column=j, value=_wert(v))
                c.font = F_NORM
                if 5 <= j <= 13:
                    c.number_format = MENGE
            sr.cell(row=i, column=4).number_format = "0.0"
            sr.cell(row=i, column=16).alignment = Alignment(wrap_text=False)
        sr.freeze_panes = "B3"
        sr.auto_filter.ref = f"A2:{get_column_letter(len(snamen))}{srlast}"
        SR = lambda col: f"'Stock Report (Info)'!${col}$3:${col}${srlast}"  # noqa: E731

    n = len(bedarf)
    last = n + 1
    namen = list(config.BEDARF_SPALTEN.values()) + [
        "Bestand verfügbar", "Zugang bis Ladedatum", "Offene Menge", "Nicht zugewiesen",
        "Tage bis Ladedatum", "Kum. Bedarf Artikel", "Fehlmenge", "Gedeckt ab", "Verspätung (Tage)", "Status"]
    L = {name: get_column_letter(i) for i, name in enumerate(namen, 1)}

    def R(name):
        return f"Positionen!${L[name]}$2:${L[name]}${last}"

    # ---------------- Bestand ----------------
    verf_arten = list(config.BESTAND_VERFUEGBAR)
    info_arten = list(config.BESTAND_INFO)
    bs["A1"] = ("Bestand je Artikel aus dem EWM-Export. Gerechnet wird nur „Verfügbar“ = "
                + " + ".join(f"{a} ({t})" for a, t in config.BESTAND_VERFUEGBAR.items())
                + ". " + ", ".join(f"{a} ({t})" for a, t in config.BESTAND_INFO.items())
                + " nur zur Info. Andere Bestandsarten werden ignoriert.")
    bs["A1"].font = F_GREY
    bnamen = (["Norm-Nr", "Kurztext"] + [f"{a} {t}" for a, t in config.BESTAND_VERFUEGBAR.items()]
              + [f"{a} {t} (Info)" for a, t in config.BESTAND_INFO.items()]
              + ["Verfügbar", "Offener Bedarf", "Bestand − Bedarf"])
    _kopf(bs, 2, bnamen, [12, 34] + [14] * (len(verf_arten) + len(info_arten)) + [13, 14, 14])
    blast = len(bestand_artikel) + 2
    k = len(verf_arten) + len(info_arten)
    c_verf, c_bed, c_diff = (get_column_letter(3 + k), get_column_letter(4 + k), get_column_letter(5 + k))
    v_last = get_column_letter(2 + len(verf_arten))
    info_col = get_column_letter(3 + len(verf_arten)) if info_arten else None
    BR = lambda col: f"Bestand!${col}$3:${col}${blast}"  # noqa: E731

    for i, b in enumerate(bestand_artikel.to_dict("records"), start=3):
        bs.cell(row=i, column=1, value=b["Norm-Nr"])
        bs.cell(row=i, column=2, value=b["Kurztext"])
        for j, a in enumerate(verf_arten + info_arten):
            c = bs.cell(row=i, column=3 + j, value=int(b[a]))
            if a in info_arten:
                c.fill = FILL_INFO
        bs[f"{c_verf}{i}"] = f"=SUM(C{i}:{v_last}{i})"
        bs[f"{c_bed}{i}"] = f"=SUMIFS({R('Offene Menge')},{R('Norm-Nr')},A{i})"
        bs[f"{c_diff}{i}"] = f"={c_verf}{i}-{c_bed}{i}"
        for col in range(1, len(bnamen) + 1):
            bs.cell(row=i, column=col).font = F_NORM
            if col >= 3:
                bs.cell(row=i, column=col).number_format = MENGE
        bs[f"{c_verf}{i}"].fill = FILL_FORMULA
    bs.freeze_panes = "C3"
    bs.auto_filter.ref = f"A2:{c_diff}{blast}"

    # ---------------- Zugänge ----------------
    zg["A1"] = ("Transporte zu uns (ZMM1000), Chargen je Lieferung zusammengefasst. Geplante Ankunft vor dem "
                "Stichtag = noch nicht empfangen -> zählt ab Stichtag (Ankunft gerechnet).")
    zg["A1"].font = F_GREY
    znamen = ["Norm-Nr", "Geplante Ankunft", "Transport", "Referenz", "Spediteur", "Lieferung",
              "Referenzbeleg", "Menge", "Ankunft gerechnet", "Kum. Zugang Artikel", "Hinweis"]
    _kopf(zg, 2, znamen, [12, 12, 11, 20, 32, 13, 13, 12, 12, 13, 18])
    zlast = max(len(zugang) + 2, 3)
    for i, z in enumerate(zugang[znamen[:8]].itertuples(index=False), start=3):
        for j, v in enumerate(z, 1):
            zg.cell(row=i, column=j, value=_wert(v)).font = F_NORM
        zg[f"I{i}"] = f"=MAX(B{i},{ST})"
        zg[f"J{i}"] = f"=H{i}" if i == 3 else f"=IF(A{i}=A{i - 1},J{i - 1}+H{i},H{i})"
        zg[f"K{i}"] = f"=IF(B{i}<{ST},\"Ankunft überfällig\",\"\")"
        for col in "IJK":
            zg[f"{col}{i}"].font, zg[f"{col}{i}"].fill = F_NORM, FILL_FORMULA
        zg[f"B{i}"].number_format = zg[f"I{i}"].number_format = DATUM
        zg[f"H{i}"].number_format = zg[f"J{i}"].number_format = MENGE
    zg.freeze_panes = "B3"
    zg.auto_filter.ref = f"A2:K{zlast}"
    zg.conditional_formatting.add(f"K3:K{zlast}", CellIsRule(operator="equal",
                                  formula=['"Ankunft überfällig"'], fill=FILL_ORA))
    ZR = lambda col: f"Zugänge!${col}$3:${col}${zlast}"  # noqa: E731

    # ---------------- Positionen ----------------
    _kopf(pos, 1, namen, [26, 13, 12, 8, 13, 12, 9, 12, 18, 30, 11, 12, 11, 10, 11,
                          11, 11, 11, 11, 9, 13, 11, 11, 10, 22])
    for i, rec in enumerate(bedarf[list(config.BEDARF_SPALTEN.values())].itertuples(index=False), start=2):
        for j, v in enumerate(rec, 1):
            pos.cell(row=i, column=j, value=_wert(v)).font = F_NORM
        r = i
        c = {name: f"{L[name]}{r}" for name in namen}
        noetig = f"({c['Kum. Bedarf Artikel']}-{c['Bestand verfügbar']})"
        f = {
            "Bestand verfügbar": f"=IFERROR(INDEX({BR(c_verf)},MATCH({c['Norm-Nr']},{BR('A')},0)),0)",
            "Zugang bis Ladedatum": f"=SUMIFS({ZR('H')},{ZR('A')},{c['Norm-Nr']},{ZR('I')},\"<=\"&{c['Ladedatum']})",
            "Offene Menge": f"={c['Bestellmenge']}-{c['Geliefert']}",
            "Nicht zugewiesen": f"=MAX(0,{c['Offene Menge']}-{c['Zugewiesen offen']})",
            "Tage bis Ladedatum": f"={c['Ladedatum']}-{ST}",
            "Kum. Bedarf Artikel": (f"={c['Offene Menge']}" if r == 2 else
                                    f"=IF({c['Norm-Nr']}={L['Norm-Nr']}{r - 1},"
                                    f"{L['Kum. Bedarf Artikel']}{r - 1}+{c['Offene Menge']},{c['Offene Menge']})"),
            "Fehlmenge": (f"=MAX(0,MIN({c['Offene Menge']},{c['Kum. Bedarf Artikel']}"
                          f"-{c['Bestand verfügbar']}-{c['Zugang bis Ladedatum']}))"),
            "Gedeckt ab": (f"=IF({c['Fehlmenge']}<=0,\"\",IF(COUNTIFS({ZR('A')},{c['Norm-Nr']},{ZR('J')},\">=\"&{noetig})=0,\"\","
                           f"_xlfn.MINIFS({ZR('I')},{ZR('A')},{c['Norm-Nr']},{ZR('J')},\">=\"&{noetig})))"),
            "Verspätung (Tage)": f"=IF({c['Gedeckt ab']}=\"\",\"\",{c['Gedeckt ab']}-{c['Ladedatum']})",
            "Status": (f"=IF({c['Offene Menge']}<=0,\"Erledigt\","
                       f"IF({c['Ladedatum']}<{ST},"
                       f"IF({c['Fehlmenge']}<=0,\"Rückstand gedeckt\","
                       f"IF({c['Gedeckt ab']}<>\"\",\"Rückstand, Zugang kommt\",\"Rückstand ungedeckt\")),"
                       f"IF({c['Fehlmenge']}<=0,\"Gedeckt\","
                       f"IF({c['Gedeckt ab']}<>\"\",\"Zugang zu spät\","
                       f"IF({c['Tage bis Ladedatum']}<={ROT},\"Engpass\",\"Fehlmenge später\")))))"),
        }
        for name, formel in f.items():
            cell = pos[c[name]]
            cell.value, cell.font, cell.fill = formel, F_NORM, FILL_FORMULA
    for name in ["Ladedatum", "Wunschtermin", "Gedeckt ab"]:
        for cell in pos[L[name]][1:]:
            cell.number_format = DATUM
    for name in ["Bestellmenge", "Geliefert", "Zugewiesen offen", "Bestand verfügbar", "Zugang bis Ladedatum",
                 "Offene Menge", "Nicht zugewiesen", "Kum. Bedarf Artikel", "Fehlmenge"]:
        for cell in pos[L[name]][1:]:
            cell.number_format = MENGE
    pos.freeze_panes = "B2"
    pos.auto_filter.ref = f"A1:{get_column_letter(len(namen))}{last}"
    _statusfarben(pos, f"{L['Status']}2:{L['Status']}{last}")

    # ---------------- Deckung je Artikel ----------------
    dnamen = ["Norm-Nr", "Kurztext", "Kunden", "Bestand verfügbar", "Gesperrt (Info)", "Zugänge unterwegs",
              "Nächster Zugang", "Zugewiesen offen", "Offene Menge gesamt", "Rückstand (Menge)",
              "Bedarf ≤ Rot-Horizont", "Bedarf ≤ Gelb-Horizont", "Fehlmenge gesamt (nach Zugängen)",
              "Erste Fehlmenge am", "Reichweite (Tage)", "Ampel"]
    dbreiten = [12, 30, 40, 12, 11, 12, 12, 12, 13, 12, 13, 13, 14, 12, 11, 9]
    if SR:
        dnamen += ["Report: Reichweite (Info)", "Report: Verantwortlich", "Report: letzte Notiz", "Notiz vom"]
        dbreiten += [12, 14, 60, 10]
    _kopf(dek, 1, dnamen, dbreiten)
    dlast = len(artikel) + 1
    for i, a in enumerate(artikel.to_dict("records"), start=2):
        r = i
        dek.cell(row=r, column=1, value=a["Norm-Nr"])
        dek.cell(row=r, column=2, value=a["Kurztext"])
        dek.cell(row=r, column=3, value=a["Kunden"][:120])
        dek[f"D{r}"] = f"=IFERROR(INDEX({BR(c_verf)},MATCH(A{r},{BR('A')},0)),0)"
        dek[f"E{r}"] = f"=IFERROR(INDEX({BR(info_col)},MATCH(A{r},{BR('A')},0)),0)" if info_col else 0
        dek[f"F{r}"] = f"=SUMIFS({ZR('H')},{ZR('A')},A{r})"
        dek[f"G{r}"] = f"=IF(F{r}=0,\"\",_xlfn.MINIFS({ZR('I')},{ZR('A')},A{r}))"
        dek[f"H{r}"] = f"=SUMIFS({R('Zugewiesen offen')},{R('Norm-Nr')},A{r})"
        dek[f"I{r}"] = f"=SUMIFS({R('Offene Menge')},{R('Norm-Nr')},A{r})"
        dek[f"J{r}"] = f"=SUMIFS({R('Offene Menge')},{R('Norm-Nr')},A{r},{R('Ladedatum')},\"<\"&{ST})"
        dek[f"K{r}"] = f"=SUMIFS({R('Offene Menge')},{R('Norm-Nr')},A{r},{R('Ladedatum')},\"<=\"&({ST}+{ROT}))"
        dek[f"L{r}"] = f"=SUMIFS({R('Offene Menge')},{R('Norm-Nr')},A{r},{R('Ladedatum')},\"<=\"&({ST}+{GELB}))"
        dek[f"M{r}"] = f"=MAX(0,I{r}-D{r}-F{r})"
        dek[f"N{r}"] = (f"=IF(COUNTIFS({R('Norm-Nr')},A{r},{R('Fehlmenge')},\">0\")=0,\"\","
                        f"_xlfn.MINIFS({R('Ladedatum')},{R('Norm-Nr')},A{r},{R('Fehlmenge')},\">0\"))")
        dek[f"O{r}"] = f"=IF(N{r}=\"\",\"\",N{r}-{ST})"
        dek[f"P{r}"] = f"=IF(N{r}=\"\",\"Grün\",IF(O{r}<={ROT},\"Rot\",IF(O{r}<={GELB},\"Gelb\",\"Grün\")))"
        for col in range(1, 17):
            cell = dek.cell(row=r, column=col)
            cell.font = F_NORM
            if col >= 4:
                cell.number_format = MENGE
        dek[f"E{r}"].fill = FILL_INFO
        dek[f"G{r}"].number_format = dek[f"N{r}"].number_format = DATUM
        dek[f"O{r}"].number_format = "0"
        if SR:
            for col, src in [("Q", "D"), ("R", "N"), ("S", "P"), ("T", "Q")]:
                text = "" if col == "Q" else "&\"\""   # Text-Spalten: leere Zellen nicht als 0 anzeigen
                dek[f"{col}{r}"] = f"=IFERROR(INDEX({SR(src)},MATCH(A{r},{SR('A')},0)){text},\"\")"
                dek[f"{col}{r}"].font, dek[f"{col}{r}"].fill = F_NORM, FILL_INFO
            dek[f"Q{r}"].number_format = "0.0"
    dek.freeze_panes = "B2"
    dek.auto_filter.ref = f"A1:{get_column_letter(len(dnamen))}{dlast}"
    for wert, fill in [("Rot", FILL_RED), ("Gelb", FILL_YEL), ("Grün", FILL_GRN)]:
        dek.conditional_formatting.add(f"P2:P{dlast}", CellIsRule(operator="equal", formula=[f'"{wert}"'], fill=fill))

    # ---------------- Bedarf je Woche ----------------
    wochen = config.WOCHEN_VORSCHAU
    _kopf(kw, 2, ["Norm-Nr", "Kurztext", "Bestand verfügbar", "Rückstand"] + [""] * wochen + ["Später"],
          [12, 30, 12, 12] + [11] * wochen + [12])
    kw["A1"] = "Offene Menge je 7-Tage-Fenster ab Stichtag (nach Ladedatum). Zeile 2: Fensterbeginn."
    kw["A1"].font = F_GREY
    for w in range(wochen):
        col = get_column_letter(5 + w)
        kw[f"{col}2"] = f"={ST}+{7 * w}"
        kw[f"{col}2"].number_format = "DD.MM."
    lastw, later = get_column_letter(4 + wochen), get_column_letter(5 + wochen)
    for i, a in enumerate(artikel.to_dict("records"), start=3):
        r = i
        kw.cell(row=r, column=1, value=a["Norm-Nr"])
        kw.cell(row=r, column=2, value=a["Kurztext"])
        kw[f"C{r}"] = f"='Deckung je Artikel'!D{r - 1}"
        kw[f"D{r}"] = f"=SUMIFS({R('Offene Menge')},{R('Norm-Nr')},$A{r},{R('Ladedatum')},\"<\"&{ST})"
        for w in range(wochen):
            col = get_column_letter(5 + w)
            kw[f"{col}{r}"] = (f"=SUMIFS({R('Offene Menge')},{R('Norm-Nr')},$A{r},"
                               f"{R('Ladedatum')},\">=\"&{col}$2,{R('Ladedatum')},\"<\"&({col}$2+7))")
        kw[f"{later}{r}"] = (f"=SUMIFS({R('Offene Menge')},{R('Norm-Nr')},$A{r},"
                             f"{R('Ladedatum')},\">=\"&({lastw}$2+7))")
        for col in range(1, 6 + wochen):
            cell = kw.cell(row=r, column=col)
            cell.font = F_NORM
            if col >= 3:
                cell.number_format = MENGE
    kw.freeze_panes = "C3"

    # ---------------- Werke & Kontakte ----------------
    kont = (bedarf.groupby(["Kunde", "Werk"])
                  .agg(ablade=("Abladestelle", lambda s: ", ".join(sorted({x for x in s if x}))[:80]),
                       we=("Warenempfänger", lambda s: ", ".join(str(x) for x in sorted(set(s)))[:80]))
                  .reset_index())
    ko["A1"] = ("Bitte die gelben Spalten einmalig pflegen. Format-Beispiel: Max Mustermann | "
                "max.mustermann@kunde.de | +49 5361 123456. Kontakte erscheinen dann im Reiter Rückstand.")
    ko["A1"].font = F_GREY
    _kopf(ko, 2, ["Schlüssel", "Kunde", "Werk", "Abladestellen", "Warenempfänger",
                  "Ansprechpartner", "E-Mail", "Telefon"], [30, 30, 8, 40, 30, 22, 30, 18])
    klast = len(kont) + 2
    for i, kk in enumerate(kont.itertuples(index=False), start=3):
        ko[f"A{i}"] = f"=B{i}&\"|\"&C{i}"
        ko[f"B{i}"], ko[f"C{i}"], ko[f"D{i}"], ko[f"E{i}"] = kk.Kunde, kk.Werk, kk.ablade, kk.we
        for col in "ABCDEFGH":
            ko[f"{col}{i}"].font = F_NORM
        for col in "FGH":
            ko[f"{col}{i}"].fill, ko[f"{col}{i}"].font = FILL_INPUT, F_INPUT
    ko.freeze_panes = "B3"
    ko.auto_filter.ref = f"A2:H{klast}"

    # ---------------- Rückstand ----------------
    st = pd.Timestamp(stichtag)
    offen = bedarf["Bestellmenge"] - bedarf["Geliefert"]
    rueck = bedarf[(bedarf["Ladedatum"] < st) & (offen > 0)].sort_values(["Ladedatum", "Kunde"])
    rk["A1"] = "Alle offenen Positionen mit Ladedatum vor dem Stichtag des Exports. Werte live aus Positionen."
    rk["A1"].font = F_GREY
    rnamen = ["Kunde", "Werk", "Abladestelle", "Lieferplan", "Norm-Nr", "Kundenmaterial", "Kurztext",
              "Ladedatum", "Wunschtermin", "Tage überfällig", "Offene Menge", "Zugewiesen offen",
              "Fehlmenge", "Gedeckt ab", "Status", "Ansprechpartner", "E-Mail"]
    rbreiten = [26, 8, 13, 12, 12, 18, 30, 11, 12, 10, 11, 11, 11, 11, 22, 22, 30]
    if SR:
        rnamen += ["Report: letzte Notiz (Info)", "Notiz vom"]
        rbreiten += [60, 10]
    _kopf(rk, 2, rnamen, rbreiten)
    rcols = get_column_letter(len(rnamen))
    KR = lambda col: f"'Werke & Kontakte'!${col}$3:${col}${klast}"  # noqa: E731
    for i, (idx, x) in enumerate(rueck.iterrows(), start=3):
        prow = idx + 2
        vals = [x["Kunde"], x["Werk"], x["Abladestelle"], x["Lieferplan"], x["Norm-Nr"],
                x["Kundenmaterial"], x["Kurztext"], _wert(x["Ladedatum"]), _wert(x["Wunschtermin"])]
        for j, v in enumerate(vals, 1):
            rk.cell(row=i, column=j, value=v)
        rk[f"J{i}"] = f"={ST}-H{i}"
        rk[f"K{i}"] = f"=Positionen!{L['Offene Menge']}{prow}"
        rk[f"L{i}"] = f"=Positionen!{L['Zugewiesen offen']}{prow}"
        rk[f"M{i}"] = f"=Positionen!{L['Fehlmenge']}{prow}"
        rk[f"N{i}"] = f"=Positionen!{L['Gedeckt ab']}{prow}"
        rk[f"O{i}"] = f"=Positionen!{L['Status']}{prow}"
        key = f"A{i}&\"|\"&B{i}"
        rk[f"P{i}"] = f"=IFERROR(INDEX({KR('F')},MATCH({key},{KR('A')},0))&\"\",\"\")"
        rk[f"Q{i}"] = f"=IFERROR(INDEX({KR('G')},MATCH({key},{KR('A')},0))&\"\",\"\")"
        if SR:
            rk[f"R{i}"] = f"=IFERROR(INDEX({SR('P')},MATCH(E{i},{SR('A')},0))&\"\",\"\")"
            rk[f"S{i}"] = f"=IFERROR(INDEX({SR('Q')},MATCH(E{i},{SR('A')},0))&\"\",\"\")"
            rk[f"R{i}"].fill = rk[f"S{i}"].fill = FILL_INFO
        for col in range(1, len(rnamen) + 1):
            rk.cell(row=i, column=col).font = F_NORM
        rk[f"H{i}"].number_format = rk[f"I{i}"].number_format = rk[f"N{i}"].number_format = DATUM
        for col in "KLM":
            rk[f"{col}{i}"].number_format = MENGE
    rlast = max(len(rueck) + 2, 3)
    rk.freeze_panes = "B3"
    rk.auto_filter.ref = f"A2:{rcols}{rlast}"
    _statusfarben(rk, f"O3:O{rlast}")

    # ---------------- Übersicht ----------------
    for col, w in zip("ABCDEF", [2, 46, 16, 34, 12, 40]):
        ue.column_dimensions[col].width = w
    ue["B1"], ue["B1"].font = "Dispo-Übersicht Lieferpläne", F_TITLE
    ue["B2"], ue["B2"].font = "Erstellt aus den jeweils neuesten SAP-Exporten.", F_GREY
    ue["B3"], ue["B3"].font = "Einstellungen", F_BOLD
    for r, label, val, fmt, hint in [
        (4, "Stichtag", datetime.combine(stichtag, datetime.min.time()), DATUM, "Datum des neuesten Bedarfsexports. Ändern = alles rechnet neu."),
        (5, "Warnhorizont Rot (Tage)", config.WARN_ROT_TAGE, "0", "Fehlmenge innerhalb dieser Tage = Rot / Engpass."),
        (6, "Warnhorizont Gelb (Tage)", config.WARN_GELB_TAGE, "0", "Fehlmenge innerhalb dieser Tage = Gelb."),
    ]:
        ue[f"B{r}"], ue[f"C{r}"], ue[f"D{r}"] = label, val, hint
        ue[f"B{r}"].font, ue[f"D{r}"].font = F_NORM, F_GREY
        ue[f"C{r}"].font, ue[f"C{r}"].fill, ue[f"C{r}"].number_format = F_INPUT, FILL_INPUT, fmt

    # Datenstand je Quelle
    zeile = 8
    ue[f"B{zeile}"], ue[f"B{zeile}"].font = "Datenstand der Quellen", F_BOLD
    zeile += 1
    for j, h in enumerate(["Quelle", "Stand", "Datei", "Alter (Tage)", "Hinweis"], start=2):
        c = ue.cell(row=zeile, column=j, value=h)
        c.font, c.fill = F_HEAD, FILL_HEAD
    for q, d, datei in datenstand:
        zeile += 1
        ue[f"B{zeile}"], ue[f"C{zeile}"], ue[f"D{zeile}"] = q, datetime.combine(d, datetime.min.time()), datei
        ue[f"E{zeile}"] = f"={ST}-C{zeile}"
        ue[f"F{zeile}"] = (f"=IF(E{zeile}>={config.QUELLE_VERALTET_TAGE},\"veraltet – neuen Export ziehen\","
                           f"IF(E{zeile}<0,\"neuer als Stichtag\",\"aktuell\"))")
        ue[f"C{zeile}"].number_format, ue[f"E{zeile}"].number_format = DATUM, "0"
        for col in "BCDEF":
            ue[f"{col}{zeile}"].font = F_NORM
            ue[f"{col}{zeile}"].border = THIN
    if datenstand:
        rng = f"F10:F{zeile}"
        ue.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"aktuell"'], fill=FILL_GRN))
        ue.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"veraltet – neuen Export ziehen"'], fill=FILL_ORA))

    PS = R("Status")
    DA = f"'Deckung je Artikel'!P2:P{dlast}"
    kpis = [
        ("Offene Einteilungen", f"=COUNTIFS({R('Offene Menge')},\">0\")", MENGE, None),
        ("Offene Menge gesamt (Stück)", f"=SUM({R('Offene Menge')})", MENGE, None),
        ("Bestand verfügbar gesamt (Stück)", f"=SUM({BR(c_verf)})", MENGE, None),
        ("Zugänge unterwegs (Stück)", f"=SUM({ZR('H')})", MENGE, None),
        ("Zugangspositionen mit überfälliger Ankunft", f"=COUNTIF({ZR('K')},\"Ankunft überfällig\")", MENGE, FILL_ORA),
        ("Rückstand ungedeckt – Positionen", f"=COUNTIF({PS},\"Rückstand ungedeckt\")", MENGE, FILL_RED),
        ("Rückstand ungedeckt – Menge", f"=SUMIFS({R('Offene Menge')},{PS},\"Rückstand ungedeckt\")", MENGE, FILL_RED),
        ("Rückstand, Zugang kommt – Positionen", f"=COUNTIF({PS},\"Rückstand, Zugang kommt\")", MENGE, FILL_ORA),
        ("Rückstand gedeckt (nur noch versenden) – Positionen", f"=COUNTIF({PS},\"Rückstand gedeckt\")", MENGE, FILL_YEL),
        ("Zugang zu spät – Positionen", f"=COUNTIF({PS},\"Zugang zu spät\")", MENGE, FILL_ORA),
        ("Engpass ohne Zugang im Rot-Horizont – Positionen", f"=COUNTIF({PS},\"Engpass\")", MENGE, FILL_RED),
        ("Engpass ohne Zugang im Rot-Horizont – Fehlmenge", f"=SUMIFS({R('Fehlmenge')},{PS},\"Engpass\")", MENGE, FILL_RED),
        ("Artikel gesamt", f"=COUNTA('Deckung je Artikel'!A2:A{dlast})", "0", None),
        ("Artikel Rot", f"=COUNTIF({DA},\"Rot\")", "0", FILL_RED),
        ("Artikel Gelb", f"=COUNTIF({DA},\"Gelb\")", "0", FILL_YEL),
        ("Artikel Grün", f"=COUNTIF({DA},\"Grün\")", "0", FILL_GRN),
        ("Artikel mit Bedarf, aber ohne Bestand", f"=COUNTIF('Deckung je Artikel'!D2:D{dlast},0)", "0", None),
    ]
    zeile += 2
    ue[f"B{zeile}"], ue[f"B{zeile}"].font = "Kennzahlen", F_BOLD
    for kk, (label, formel, fmt, fill) in enumerate(kpis, start=zeile + 1):
        ue[f"B{kk}"], ue[f"C{kk}"] = label, formel
        ue[f"B{kk}"].font = F_NORM
        ue[f"C{kk}"].font, ue[f"C{kk}"].number_format = F_BOLD, fmt
        ue[f"B{kk}"].border = ue[f"C{kk}"].border = THIN
        if fill:
            ue.conditional_formatting.add(f"C{kk}", CellIsRule(operator="greaterThan", formula=["0"], fill=fill))
    zeile += len(kpis)

    regeln = [
        "So wird gerechnet",
        "Offene Menge = Bestellmenge − Geliefert. Nicht zugewiesen = Offene Menge − Zugewiesen offen (= SAP „Kalan“).",
        "Bestand verfügbar = " + " + ".join(f"{a} ({t})" for a, t in config.BESTAND_VERFUEGBAR.items())
        + " aus dem EWM-Export. " + ", ".join(f"{a} ({t})" for a, t in config.BESTAND_INFO.items())
        + " nur Info, alle anderen Bestandsarten ignoriert.",
        "Zugänge (ZMM1000): zählen ab geplanter Ankunft bei uns; überfällige, nicht empfangene Ankünfte ab Stichtag.",
        "Deckung: Je Artikel werden Bestand + Zugänge, die bis zum Ladedatum da sind, in Reihenfolge des Ladedatums "
        "auf die offenen Einteilungen verteilt. Was fehlt, ist die Fehlmenge.",
        "Gedeckt ab = Ankunft des Zugangs, mit dem die Fehlmenge gedeckt wäre. Verspätung = Gedeckt ab − Ladedatum.",
        "Status: Rückstand ungedeckt / Rückstand, Zugang kommt / Rückstand gedeckt (nur versenden) / "
        "Zugang zu spät / Engpass (kein Zugang, im Rot-Horizont) / Fehlmenge später / Gedeckt / Erledigt.",
        "Ampel: Rot = erste Fehlmenge im Rot-Horizont oder überfällig, Gelb = im Gelb-Horizont, Grün = später/nie.",
        "Stock Report des Versenders: nur Nebeninfo (Reichweite, Notizen) – geht nicht in die Rechnung ein.",
        "Es wird immer die neueste Datei je Quelle verwendet; ältere Quellen sind oben als veraltet markiert.",
        "Gelb hinterlegte, blaue Werte = Eingaben. Hellblau = Formeln. Grau = nur Info.",
    ]
    start = zeile + 3
    for kk, t in enumerate(regeln, start=start):
        ue[f"B{kk}"] = t
        ue[f"B{kk}"].font = F_BOLD if kk == start else F_NORM

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title != "Übersicht"
    wb.save(ziel)
